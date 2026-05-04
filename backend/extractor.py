from __future__ import annotations

import csv
import ipaddress
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

import openpyxl
import pdfplumber
from docx import Document

EMAIL_RE = re.compile(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
URL_RE = re.compile(r"https?://[^\s\]\[\)\(<>\"']+")
DOMAIN_RE = re.compile(r"\b(?:[A-Za-z0-9-]+\.)+[A-Za-z]{2,63}\b")
CVE_RE = re.compile(r"\bCVE-\d{4}-\d{4,7}\b", re.I)
ASN_RE = re.compile(r"\b(?:AS|ASN\s*)\d{1,10}\b", re.I)
HASH_RE = re.compile(r"\b(?:[A-Fa-f0-9]{32}|[A-Fa-f0-9]{40}|[A-Fa-f0-9]{64})\b")
CIDR_RE = re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}/\d{1,2}\b|\b[0-9a-fA-F:]+/\d{1,3}\b")
IPV4_RE = re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}\b")
IPV6_RE = re.compile(r"\b(?:[0-9a-fA-F]{1,4}:){2,}[0-9a-fA-F:]{1,4}\b")


def _iter_source(input_data) -> Iterable[tuple[str, str]]:
    if isinstance(input_data, str):
        p = Path(input_data)
        if p.exists():
            if p.suffix in {".txt", ".log", ".csv"}:
                yield p.name, p.read_text(errors="ignore")
            elif p.suffix in {".xlsx", ".xls"}:
                wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
                for sh in wb.worksheets:
                    for row in sh.iter_rows(values_only=True):
                        yield p.name, " ".join(str(c) for c in row if c is not None)
            elif p.suffix == ".docx":
                doc = Document(str(p))
                for para in doc.paragraphs:
                    yield p.name, para.text
                core = doc.core_properties
                for k in ["author", "last_modified_by", "created", "modified", "comments", "category"]:
                    v = getattr(core, k, None)
                    if v:
                        yield p.name, f"METADATA::{k}::{v}"
            elif p.suffix == ".pdf":
                with pdfplumber.open(str(p)) as pdf:
                    for pg in pdf.pages:
                        yield p.name, pg.extract_text() or ""
                    for k, v in (pdf.metadata or {}).items():
                        if v:
                            yield p.name, f"METADATA::{k}::{v}"
        else:
            yield "raw_text", input_data


def extract(input_data):
    seen, out = {}, []
    for source_file, blob in _iter_source(input_data):
        for line_no, line in enumerate(blob.splitlines() or [blob], 1):
            tokens = []
            if line.startswith("METADATA::"):
                tokens.append(("Metadata", line))
            for rtype, rgx in [("IPv4", IPV4_RE), ("IPv6", IPV6_RE), ("Email", EMAIL_RE), ("Domain", DOMAIN_RE), ("URL", URL_RE),
                              ("CIDR", CIDR_RE), ("ASN", ASN_RE), ("Hash", HASH_RE), ("CVE", CVE_RE)]:
                tokens += [(rtype, m.group(0)) for m in rgx.finditer(line)]
            for typ, val in tokens:
                norm = val.lower()
                notes = []
                if typ == "IPv4":
                    try:
                        ip = ipaddress.ip_address(val)
                        if ip.is_private:
                            notes.append("Private RFC1918")
                    except Exception:
                        continue
                if typ == "IPv6":
                    try:
                        norm = ipaddress.ip_address(val).exploded.lower()
                    except Exception:
                        continue
                key = f"{typ}::{norm}"
                if key in seen:
                    continue
                seen[key] = True
                out.append({"ioc_id": str(uuid.uuid4()), "type": typ, "value": val, "normalized": norm, "source_file": source_file,
                            "source_line": line_no, "context": line[:160], "extraction_notes": notes,
                            "pre_flags": {"is_private_ip": "Private RFC1918" in notes, "format_valid": True},
                            "timestamp": datetime.now(timezone.utc).isoformat()})
    return out
